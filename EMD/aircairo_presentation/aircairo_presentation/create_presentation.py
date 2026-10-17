from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_aircairo_presentation(filename):
    prs = Presentation()
    
    # Define colors
    aircairo_blue = RGBColor(0, 85, 165) # Dark blue
    aircairo_light_blue = RGBColor(0, 169, 224) # Light blue
    accent_orange = RGBColor(255, 153, 0) # Orange
    white = RGBColor(255, 255, 255)
    light_gray = RGBColor(240, 240, 240)

    # Helper function to set slide background
    def set_slide_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Slide 1: Title Slide (Dark Blue)
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    set_slide_bg(slide1, aircairo_blue)
    
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    
    title.text = "AirCairo Seatmap Redesign"
    title.text_frame.paragraphs[0].font.color.rgb = white
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "Project Kick-off Meeting\n[Date]"
    subtitle.text_frame.paragraphs[0].font.color.rgb = aircairo_light_blue
    
    # Slide 2: Vision / Goal
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    set_slide_bg(slide2, white)
    
    title2 = slide2.shapes.title
    title2.text = "الرؤية والهدف (Vision & Goals)"
    title2.text_frame.paragraphs[0].font.color.rgb = aircairo_blue
    
    body_shape2 = slide2.shapes.placeholders[1]
    tf = body_shape2.text_frame
    tf.text = "لماذا نقوم بإعادة تصميم خريطة المقاعد؟"
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.color.rgb = aircairo_blue
    
    p = tf.add_paragraph()
    p.text = "تحسين تجربة المستخدم (UX) وجعلها أكثر سلاسة وسرعة."
    p.font.size = Pt(24)
    
    p = tf.add_paragraph()
    p.text = "زيادة وضوح المقاعد المميزة (Extra Legroom) لزيادة المبيعات."
    p.font.size = Pt(24)

    p = tf.add_paragraph()
    p.text = "توفير تصميم متجاوب (Mobile-Friendly) لجميع الأجهزة."
    p.font.size = Pt(24)
    
    # Slide 3: The New Design (Mockup Placeholder)
    slide3 = prs.slides.add_slide(prs.slide_layouts[5]) # Title only layout
    set_slide_bg(slide3, light_gray)
    
    title3 = slide3.shapes.title
    title3.text = "التصميم الجديد (The New Seatmap)"
    title3.text_frame.paragraphs[0].font.color.rgb = aircairo_blue
    title3.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Add a placeholder shape for the image
    left = Inches(2)
    top = Inches(2)
    width = Inches(6)
    height = Inches(4.5)
    shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = white
    shape.line.color.rgb = aircairo_light_blue
    
    tf3 = shape.text_frame
    tf3.text = "[صورة التصميم الجديد للـ Seatmap تضاف هنا]"
    tf3.paragraphs[0].font.color.rgb = aircairo_blue
    tf3.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Slide 4: Color Legend
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    set_slide_bg(slide4, white)
    
    title4 = slide4.shapes.title
    title4.text = "دلالات الألوان (Color Coding)"
    title4.text_frame.paragraphs[0].font.color.rgb = aircairo_blue
    
    body_shape4 = slide4.shapes.placeholders[1]
    tf4 = body_shape4.text_frame
    tf4.text = "تم اختيار الألوان لتكون واضحة ومريحة للعين:"
    
    p = tf4.add_paragraph()
    p.text = "الأزرق الفاتح / الأخضر: مقعد متاح."
    p.font.color.rgb = aircairo_light_blue
    
    p = tf4.add_paragraph()
    p.text = "البرتقالي: مقعد بمساحة إضافية (Extra Legroom)."
    p.font.color.rgb = accent_orange
    
    p = tf4.add_paragraph()
    p.text = "الرمادي: مقعد محجوز أو غير متاح."
    p.font.color.rgb = RGBColor(128, 128, 128)

    # Slide 5: Roadmap
    slide5 = prs.slides.add_slide(bullet_slide_layout)
    set_slide_bg(slide5, white)
    
    title5 = slide5.shapes.title
    title5.text = "خريطة الطريق (Timeline)"
    title5.text_frame.paragraphs[0].font.color.rgb = aircairo_blue
    
    body_shape5 = slide5.shapes.placeholders[1]
    tf5 = body_shape5.text_frame
    tf5.text = "المراحل القادمة للمشروع:"
    
    p = tf5.add_paragraph()
    p.text = "الأسبوع الأول: الانتهاء من التصميم النهائي (UI/UX)."
    
    p = tf5.add_paragraph()
    p.text = "الأسبوع الثاني: بدء عملية البرمجة (Development)."
    
    p = tf5.add_paragraph()
    p.text = "الأسبوع الثالث: الاختبار (Testing) وإطلاق النسخة التجريبية."

    # Slide 6: Q&A
    slide6 = prs.slides.add_slide(title_slide_layout)
    set_slide_bg(slide6, aircairo_blue)
    
    title6 = slide6.shapes.title
    title6.text = "شكراً لكم!\nالأسئلة والنقاش؟"
    title6.text_frame.paragraphs[0].font.color.rgb = white
    
    prs.save(filename)
    print(f"Presentation saved to {filename}")

if __name__ == '__main__':
    create_aircairo_presentation('AirCairo_Seatmap_Kickoff.pptx')
